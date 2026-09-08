from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from src.utils import ScheduleItem, combine_date_and_time


SCHEDULE_SHEET = "Class Schedule PGPEx"
TIME_RANGE = re.compile(r"(\d{1,2}:\d{2})\s*-\s*(\d{1,2}:\d{2})\s*hrs", re.IGNORECASE)


def parse_excel_schedule(path: Path) -> list[ScheduleItem]:
    """Read the current timetable by cell position, excluding hidden legacy sheets."""
    workbook = load_workbook(path, data_only=True)
    try:
        if SCHEDULE_SHEET not in workbook.sheetnames:
            raise ValueError(f"Missing timetable sheet: {SCHEDULE_SHEET}")
        sheet = workbook[SCHEDULE_SHEET]
        if sheet.sheet_state != "visible":
            raise ValueError(f"Timetable sheet is hidden: {SCHEDULE_SHEET}")

        slots = {}
        header_row = None
        for row in sheet.iter_rows():
            for cell in row:
                match = TIME_RANGE.fullmatch(str(cell.value or "").strip())
                if match:
                    slots[cell.column] = tuple(value.zfill(5) for value in match.groups())
                    header_row = cell.row
            if slots:
                break
        if header_row is None:
            raise ValueError("No timetable time-slot headers found")

        items = []
        for row in sheet.iter_rows(min_row=header_row + 1):
            if any(str(cell.value or "").strip() == "Course Name" for cell in row):
                break
            day = sheet.cell(row[0].row, 2).value
            if isinstance(day, datetime):
                day = day.date()
            if not isinstance(day, date):
                continue

            for column, (start_time, end_time) in slots.items():
                cell = sheet.cell(row[0].row, column)
                if cell.data_type == "e":
                    raise ValueError(f"Spreadsheet error at {SCHEDULE_SHEET}!{cell.coordinate}: {cell.value}")
                value = " ".join(str(cell.value or "").split())
                if not value:
                    continue
                merged = next((area for area in sheet.merged_cells.ranges if cell.coordinate in area), None)
                if merged is not None and merged.min_col <= min(slots) and merged.max_col >= max(slots):
                    last_day = sheet.cell(merged.max_row, 2).value
                    if isinstance(last_day, datetime):
                        last_day = last_day.date()
                    if not isinstance(last_day, date):
                        raise ValueError(f"Missing end date for merged block {merged}")
                    items.append(ScheduleItem(value, day, last_day + timedelta(days=1)))
                    continue

                # Slashes separate parallel subjects; empty placeholders create no events.
                for title in (part.strip() for part in value.split("/")):
                    if title:
                        items.append(ScheduleItem(
                            title,
                            combine_date_and_time(day, start_time),
                            combine_date_and_time(day, end_time),
                        ))
        return items
    finally:
        workbook.close()
